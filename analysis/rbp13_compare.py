import os

import numpy as np
import pandas as pd
from joblib import Parallel, delayed
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

import rbp
from backtest import HORIZON, base_rate, load_panel, predictors, run_backtest, score
from logistic import logistic
from ml import HISTORY_START, VALIDATION, neural_net, walk_forward

# RBP, the logistic regression and the neural net on the same thirteen predictors (analysis/rbp.py)
# and the same forecast dates, 1991-01 to 2025-12, in one efficacy table for Kritzman.
# Written by Claude Code at Flynn's request (2026-10-20).
#
# RBP: data/rbp/rbp13_forecasts.csv (run analysis/rbp.py first), all forecasts and the top 50% and
# 25% by fit (rbp.high_fit). Logistic regression: analysis/logistic.py, as in the trading rule,
# trained on the months with every predictor (from 1957, when HHI starts). Neural net: analysis/ml.py,
# trained from 1930-10 with missing values filled; its penalty is chosen on 1961-1990 by Brier score
# before the test run, as in ml.py.
#
# Rows per threshold, as in data/rbp/rbp_results_1991_2026.xlsx, but the last row is replaced by two
# (Kritzman, 2026-10-20): the share of loss episodes warned of before the fall began, (1) by a warning
# at the episode's onset, its first loss month, and (2) by a warning at any of the three month-ends
# up to and including the onset. An episode: loss months less than three months apart. An episode
# counts in a column's sample if its onset date is in it.
#
# Writes data/rbp/rbp13_compare_forecasts.csv and data/rbp/rbp13_compare_1991_2025.xlsx.
# Run from the repository root: python analysis/rbp13_compare.py

FIRST, LAST = "1991-01-31", "2025-12-31"
THRESHOLDS = [0.10, 0.15, 0.20]


def panel_13():
    """The panel with RBP's thirteen predictors, built as analysis/rbp.py builds them."""
    panel = rbp.total_return_inputs(rbp.total_return_target(load_panel()))
    return rbp.concentration(rbp.sentiment(rbp.cape(rbp.yield_curve_long(rbp.credit_spread(rbp.fragility(panel))))))


def neural_net_forecasts(panel, X_all):
    """Choose the net's penalty on 1961-1990, then run it once on the test window."""
    reference = walk_forward(base_rate, X_all, panel, *VALIDATION)
    alphas = (0.1, 1.0, 10.0)
    runs = Parallel(n_jobs=len(alphas))(delayed(walk_forward)(neural_net(a), X_all, panel, *VALIDATION) for a in alphas)
    briers = {a: score(r, reference)["Brier"] for a, r in zip(alphas, runs)}
    chosen = min(briers, key=briers.get)
    print("Neural net, validation Brier by alpha:", {a: round(b, 4) for a, b in briers.items()}, "chosen:", chosen)
    return walk_forward(neural_net(chosen), X_all, panel, FIRST, LAST)["Probability"], chosen


def onsets(y):
    """Onset date of each loss episode: loss months less than three months apart are one episode."""
    out, prev = [], None
    for d in y.index[y == 1]:
        if prev is None or (d.year - prev.year) * 12 + d.month - prev.month >= 3:
            out.append(d)
        prev = d
    return out


def rows(r, t, onset_dates, all_dates):
    """The table's rows for one column's forecasts r at threshold t."""
    w = r[r["Probability"] >= t]
    warned = r["Probability"] >= t
    in_sample = [d for d in onset_dates if d in r.index]
    at_onset = sum(bool(warned[d]) for d in in_sample)
    up_to = 0
    for d in in_sample:
        window = all_dates[max(0, all_dates.get_loc(d) - (HORIZON - 1)): all_dates.get_loc(d) + 1]
        up_to += bool(warned.reindex(window, fill_value=False).any())
    n = len(in_sample)
    share = lambda k: f"{k} of {n} ({k / n:.0%})" if n else "0 of 0"
    return [
        len(w),
        w["Return_3m"].mean(),
        r.loc[~warned, "Return_3m"].mean(),
        int(len(w) - w["Event_A"].sum()),
        w["Event_A"].mean() if len(w) else np.nan,
        share(at_onset),
        share(up_to),
    ]


METRICS = ["Number of warnings", "Avg 3-month return after warning", "Avg 3-month return without warning",
           "False positives", "Fraction followed by a 10% loss",
           "Episodes warned at onset (before the fall)", "Episodes warned in the 3 months up to onset"]
FORMATS = ["0", "0.0%", "0.0%", "0", "0.0%", "@", "@"]


def write_xlsx(columns, path, note):
    head_fill, band_fill, bold = PatternFill("solid", fgColor="DDE4EE"), PatternFill("solid", fgColor="F3F5F8"), Font(bold=True)
    all_dates = columns["Full RBP"].index
    onset_dates = onsets(columns["Full RBP"]["Event_A"])
    wb = Workbook()
    ws = wb.active
    ws.title = "Results"
    ws.append(["Threshold", "Metric", *columns])
    for c in ws[1]:
        c.font, c.fill = bold, head_fill
    for i, t in enumerate(THRESHOLDS):
        vals = {name: rows(r, t, onset_dates, all_dates) for name, r in columns.items()}
        start = ws.max_row + 1
        for j, m in enumerate(METRICS):
            ws.append([f"{t:.0%}" if j == 0 else None, m, *(vals[n][j] for n in columns)])
            for c in ws[ws.max_row][2:]:
                c.number_format = FORMATS[j]
            if i == 1:
                for c in ws[ws.max_row]:
                    c.fill = band_fill
        ws.merge_cells(start_row=start, end_row=start + len(METRICS) - 1, start_column=1, end_column=1)
        ws.cell(start, 1).font = bold
        ws.cell(start, 1).alignment = Alignment(vertical="center")
    ws.column_dimensions["A"].width = 11
    ws.column_dimensions["B"].width = 46
    for k in range(len(columns)):
        ws.column_dimensions[chr(ord("C") + k)].width = 20

    descriptions = {
        "logret_past_3m_tr": ("Past 3-month return", "S&P 500 total return over the previous three months"),
        "vol_trailing63_tr": ("Trailing volatility", "63-day volatility of S&P 500 total returns"),
        "ip_growth_12m": ("Industrial production growth", "12-month growth, as first published"),
        "cpi_infl_12m": ("Inflation", "12-month CPI inflation, as first published"),
        "unrate_chg_12m": ("Unemployment change", "12-month change in the unemployment rate, as first published"),
        "yc_long": ("Yield-curve slope", "10-year Treasury yield minus 3-month bill rate"),
        "turb17": ("Turbulence", "How unusual industry returns are relative to their history"),
        "ar17_shift": ("Absorption-ratio shift", "How tightly stocks move together, relative to the past year"),
        "credit_baa_aaa": ("Credit spread", "Baa minus Aaa corporate bond yield"),
        "cape": ("CAPE", "Shiller's cyclically adjusted P/E"),
        "sentiment_umich": ("Consumer sentiment", "University of Michigan consumer sentiment index"),
        "hhi": ("Concentration (HHI)", "Sum of squared market-cap weights of S&P 500 members"),
        "gjr_skewt_vol_q": ("GARCH volatility", "GJR-GARCH (skewed-t) forecast of volatility over the next three months"),
    }
    ps = wb.create_sheet("Predictors")
    ps.append(["Predictor", "Description"])
    for c in ps[1]:
        c.font, c.fill = bold, head_fill
    for p in rbp.PREDICTORS:
        ps.append(list(descriptions[p]))
    ps.column_dimensions["A"].width = 30
    ps.column_dimensions["B"].width = 62

    ns = wb.create_sheet("Notes")
    for line in note:
        ns.append([line])
    ns.column_dimensions["A"].width = 120
    wb.save(path)


if __name__ == "__main__":
    panel = panel_13()
    X = predictors(panel, rbp.PREDICTORS)
    X_all = X.loc[HISTORY_START:]

    f = pd.read_csv("data/rbp/rbp13_forecasts.csv", parse_dates=["Date"]).set_index("Date").loc[FIRST:LAST]
    logit = run_backtest(logistic, X, panel, FIRST, LAST)["Probability"]
    net, alpha = neural_net_forecasts(panel, X_all)

    both = pd.DataFrame({"RBP": f["Probability"], "Fit": f["Fit"], "Logistic regression": logit,
                         "Neural net": net}).join(panel[["Event_A", "Return_3m"]])
    assert both.notna().all(axis=None), "every model must have a forecast on every date"
    os.makedirs("data/rbp", exist_ok=True)
    both.to_csv("data/rbp/rbp13_compare_forecasts.csv", float_format="%.6f")

    col = lambda name, idx=both.index: both.loc[idx, [name, "Event_A", "Return_3m"]].rename(columns={name: "Probability"})
    columns = {
        "Full RBP": col("RBP"),
        "RBP, top 50% by fit": col("RBP", both.index[rbp.high_fit(both["Fit"], 0.50)]),
        "RBP, top 25% by fit": col("RBP", both.index[rbp.high_fit(both["Fit"], 0.25)]),
        "Logistic regression": col("Logistic regression"),
        "Neural net": col("Neural net"),
    }
    n_onsets = len(onsets(both["Event_A"]))
    note = [
        f"Monthly forecasts {both.index[0]:%b %Y} to {both.index[-1]:%b %Y} ({len(both)} forecasts, "
        f"{int(both['Event_A'].sum())} loss months in {n_onsets} episodes), the same dates and the same "
        f"{len(rbp.PREDICTORS)} predictors for every model. A warning is a forecast at or above the threshold.",
        "Target: S&P 500 total return over the next three months of -10% or worse.",
        "Top 50% / 25% by fit: RBP forecasts whose fit is at or above that quantile of the previous 60 forecasts' fits.",
        "Loss episode: loss months less than three months apart. Onset: its first loss month, the month-end before the fall.",
        "Episodes warned at onset: the forecast at the onset was a warning. In the 3 months up to onset: a warning at the "
        "onset or at either of the two month-ends before it. Counted over the episodes whose onset is in each column's sample.",
        "Logistic regression: standardized, trained on months with every predictor (from 1957). "
        f"Neural net: one hidden layer of 8 units, L2 penalty {alpha} chosen on 1961-1990, average of 5 seeds, trained from 1930.",
    ]
    write_xlsx(columns, "data/rbp/rbp13_compare_1991_2025.xlsx", note)

    pd.set_option("display.width", 250)
    for name, r in columns.items():
        s = score(r)
        print(f"{name:22s} forecasts {s['Forecasts']}  AUC {s['AUC']:.3f}  mean prob {s['Mean prob']:.3f}")
